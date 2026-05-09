"""S3 이미지 로컬 다운로드 스크립트."""
import csv
import json
import os
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from urllib.parse import urlparse

import openpyxl
import requests

BASE_DIR = Path(__file__).parent
UPLOAD_DIR = BASE_DIR / "uploads" / "checklist"
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

CSV_FILE = Path("/Users/Kang/Desktop/fuckbusan/제목 없음/공공디자인 진단 플랫폼 데이터시트_양식 (1).csv")
XLSX_GENERAL = Path("/Users/Kang/Desktop/fuckbusan/제목 없음/2차진단_일반인_251210.xlsx")
XLSX_EXPERT = Path("/Users/Kang/Desktop/fuckbusan/제목 없음/2차진단_전문가_251210.xlsx")
MAP_OUT = BASE_DIR / "image_url_map.json"


def collect_urls() -> set[str]:
    urls = set()
    # CSV
    with open(CSV_FILE, encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            u = (row.get("이미지경로") or "").strip()
            if u:
                urls.add(u)
    # XLSX files
    for path in [XLSX_GENERAL, XLSX_EXPERT]:
        wb = openpyxl.load_workbook(path, read_only=True)
        ws = wb.active
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        try:
            img_col = headers.index("이미지경로") + 1
        except ValueError:
            continue
        for row in ws.iter_rows(min_row=2, values_only=True):
            u = str(row[img_col - 1] or "").strip()
            if u and u != "None":
                urls.add(u)
        wb.close()
    return urls


def download_one(url: str) -> tuple[str, str | None]:
    parsed = urlparse(url)
    filename = Path(parsed.path).name
    if not Path(filename).suffix:
        filename += ".jpg"
    dest = UPLOAD_DIR / filename
    if dest.exists():
        return url, f"/uploads/checklist/{filename}"
    try:
        r = requests.get(url, timeout=15)
        r.raise_for_status()
        dest.write_bytes(r.content)
        return url, f"/uploads/checklist/{filename}"
    except Exception as e:
        print(f"  FAIL {url}: {e}")
        return url, None


def main():
    urls = collect_urls()
    print(f"Collected {len(urls)} unique image URLs")

    mapping: dict[str, str | None] = {}
    success = 0
    fail = 0

    with ThreadPoolExecutor(max_workers=10) as pool:
        futures = {pool.submit(download_one, u): u for u in urls}
        for i, fut in enumerate(as_completed(futures), 1):
            url, local = fut.result()
            mapping[url] = local
            if local:
                success += 1
            else:
                fail += 1
            if i % 30 == 0:
                print(f"  {i}/{len(urls)} done...")

    MAP_OUT.write_text(json.dumps(mapping, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\nDone: {success} success, {fail} fail")
    print(f"Mapping saved → {MAP_OUT}")
    print(f"Images saved → {UPLOAD_DIR}")


if __name__ == "__main__":
    main()
