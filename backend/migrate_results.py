"""Migrate diagnosis response data into checklist_result table."""
import csv
import json
import pymysql
import openpyxl
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor

DB = dict(host='127.0.0.1', port=3306, user='busan', password='busanpw',
          db='busan_design_db', charset='utf8mb4')

IMAGE_MAP_PATH = '/Users/Kang/Desktop/fuckbusan/backend/image_url_map.json'
CSV_PATH = '/Users/Kang/Desktop/fuckbusan/제목 없음/공공디자인 진단 플랫폼 데이터시트_양식 (1).csv'
XLSX_GENERAL = '/Users/Kang/Desktop/fuckbusan/제목 없음/2차진단_일반인_251210.xlsx'
XLSX_EXPERT  = '/Users/Kang/Desktop/fuckbusan/제목 없음/2차진단_전문가_251210.xlsx'

TEXT_SCORES = {'적합', '부적합', '해당없음', '해당 없음'}

def load_image_map():
    with open(IMAGE_MAP_PATH, encoding='utf-8') as f:
        return json.load(f)

def load_user_map(conn):
    with conn.cursor() as cur:
        cur.execute("SELECT user_id, ID FROM users")
        return {row['ID']: row['user_id'] for row in cur.fetchall()}

def parse_dt(s):
    if not s:
        return None
    s = str(s).strip()
    for fmt in ('%Y.%m.%d %H:%M:%S', '%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M:%S.%f'):
        try:
            return datetime.strptime(s[:19], fmt[:len(s[:19].replace('.', '-').replace(':', ':'))])
        except Exception:
            pass
    for fmt in ('%Y.%m.%d %H:%M:%S', '%Y-%m-%d %H:%M:%S'):
        try:
            return datetime.strptime(s[:19], fmt)
        except Exception:
            pass
    return None

def parse_score(val):
    if val is None:
        return None, None
    s = str(val).strip()
    if s in TEXT_SCORES:
        return None, s
    try:
        return int(float(s)), None
    except Exception:
        return None, s if s else None

def map_image(url, image_map):
    if not url:
        return None, False
    url = str(url).strip()
    if not url or url == 'None':
        return None, False
    local = image_map.get(url)
    if local:
        return local, True
    return url, False

def row_to_tuple(r, user_map, image_map):
    score, mansok = parse_score(r.get('점수'))
    img, mapped = map_image(r.get('이미지경로'), image_map)
    uid = user_map.get(str(r.get('ID') or '').strip())
    lat = r.get('위도')
    lng = r.get('경도')
    try:
        lat = float(lat) if lat is not None and str(lat).strip() else None
    except Exception:
        lat = None
    try:
        lng = float(lng) if lng is not None and str(lng).strip() else None
    except Exception:
        lng = None
    return (
        str(r.get('진단지역') or '')[:100] or None,
        uid,
        str(r.get('ID') or '')[:50] or None,
        parse_dt(r.get('등록일시')),
        None,  # district_code
        lat, lng,
        str(r.get('대분류') or '')[:50] or None,
        str(r.get('중분류') or '')[:50] or None,
        str(r.get('질문_기준') or '')[:255] or None,
        str(r.get('질문_내용') or '')[:4000] or None,
        score,
        str(r.get('리뷰') or '')[:2000] or None,
        mansok,
        str(img)[:500] if img else None,
        mapped
    )

def read_csv(path):
    with open(path, encoding='utf-8-sig') as f:
        return list(csv.DictReader(f))

def read_xlsx(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    rows = []
    for r in range(2, ws.max_row + 1):
        row = {headers[c-1]: ws.cell(r, c).value for c in range(1, ws.max_column + 1)}
        rows.append(row)
    wb.close()
    return rows

SQL = """
INSERT INTO checklist_result
  (진단지역, user_id, ID, created_at, district_code, 위도, 경도,
   대분류, 중분류, 질문기준, answers, 점수, 리뷰, 만족도, 이미지경로)
VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
"""

def batch_insert(conn, tuples, source_name, errors):
    inserted = 0
    img_replaced = 0
    data = []
    mapped_flags = []
    for t in tuples:
        *fields, was_mapped = t
        data.append(tuple(fields))
        mapped_flags.append(was_mapped)

    BATCH = 500
    for i in range(0, len(data), BATCH):
        chunk = data[i:i+BATCH]
        chunk_flags = mapped_flags[i:i+BATCH]
        try:
            with conn.cursor() as cur:
                cur.executemany(SQL, chunk)
            conn.commit()
            inserted += len(chunk)
            img_replaced += sum(chunk_flags)
        except Exception as e:
            conn.rollback()
            print(f"  [ERROR] batch {i//BATCH} in {source_name}: {e}")
            errors[0] += len(chunk)
    return inserted, img_replaced

def main():
    image_map = load_image_map()
    print(f"Image map loaded: {len(image_map)} entries")

    conn = pymysql.connect(**DB, cursorclass=pymysql.cursors.DictCursor)
    user_map = load_user_map(conn)
    print(f"User map loaded: {len(user_map)} users")

    with conn.cursor() as cur:
        cur.execute("DELETE FROM checklist_result")
    conn.commit()
    print("Cleared checklist_result table")

    errors = [0]
    total = 0
    total_img = 0

    # File 1: CSV
    print("\n[1/3] CSV 일반 데이터...")
    rows1 = read_csv(CSV_PATH)
    tuples1 = []
    for r in rows1:
        try:
            tuples1.append(row_to_tuple(r, user_map, image_map))
        except Exception as e:
            print(f"  row parse error: {e}")
            errors[0] += 1
    n, ni = batch_insert(conn, tuples1, 'CSV', errors)
    print(f"  → inserted {n}, image replaced {ni}")
    total += n; total_img += ni

    # File 2: xlsx 일반인
    print("\n[2/3] xlsx 일반인...")
    rows2 = read_xlsx(XLSX_GENERAL)
    tuples2 = []
    for r in rows2:
        try:
            tuples2.append(row_to_tuple(r, user_map, image_map))
        except Exception as e:
            print(f"  row parse error: {e}")
            errors[0] += 1
    n, ni = batch_insert(conn, tuples2, 'xlsx_general', errors)
    print(f"  → inserted {n}, image replaced {ni}")
    total += n; total_img += ni

    # File 3: xlsx 전문가
    print("\n[3/3] xlsx 전문가...")
    rows3 = read_xlsx(XLSX_EXPERT)
    tuples3 = []
    for r in rows3:
        try:
            tuples3.append(row_to_tuple(r, user_map, image_map))
        except Exception as e:
            print(f"  row parse error: {e}")
            errors[0] += 1
    n, ni = batch_insert(conn, tuples3, 'xlsx_expert', errors)
    print(f"  → inserted {n}, image replaced {ni}")
    total += n; total_img += ni

    with conn.cursor() as cur:
        cur.execute("SELECT COUNT(*) as cnt FROM checklist_result")
        db_count = cur.fetchone()['cnt']

    conn.close()

    print(f"""
=== DONE ===
CSV 일반:       {len(tuples1)} rows
xlsx 일반인:    {len(tuples2)} rows
xlsx 전문가:    {len(tuples3)} rows
총 insert:      {total}
DB count:       {db_count}
이미지 로컬교체: {total_img}
에러:           {errors[0]}
""")

if __name__ == '__main__':
    main()
